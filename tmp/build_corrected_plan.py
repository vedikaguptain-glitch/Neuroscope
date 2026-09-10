from datetime import date
from pathlib import Path

from docx import Document

from build_simple_protocol import (
    add_hyperlink,
    bullet,
    configure_styles,
    heading,
    numbered,
    paragraph,
    set_font,
)


ROOT = Path(r"C:\Users\USER\Desktop\code\Neuroscope")
OUT = ROOT / "output" / "documents" / "neuroscope-data-collection-plan-corrected.docx"


def add_title(doc):
    p = doc.add_paragraph()
    r = p.add_run("NEUROSCOPE Data Collection Plan")
    set_font(r, size=20, bold=True)

    p = doc.add_paragraph()
    r = p.add_run("Critique of the earlier plan and a corrected plan for the science fair")
    set_font(r, italic=True)

    p = doc.add_paragraph()
    r = p.add_run("Prepared: ")
    set_font(r, bold=True)
    r = p.add_run(date.today().strftime("%d %B %Y"))
    set_font(r)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "NEUROSCOPE Data Collection Plan"
    doc.core_properties.author = "NEUROSCOPE research team"
    add_title(doc)

    heading(doc, "Final decision")
    paragraph(doc, "Use whole class sections as the student sample. Randomly choose one section from each participating grade, invite every student in those sections who has the required permission, and open the normal app link during planned sessions. This uses the ordinary link and avoids the large amount of work involved in messaging randomly chosen students one by one.")
    paragraph(doc, "Invite all eligible teachers, but keep their results separate. The main result should be about students from the selected classes in this school. Teacher results should be described as a small adult comparison, not as proof that the model works for adults.")
    for text in [
        "Do not post the link across the whole school. Post or send it only to the selected classes and the teacher group.",
        "Collect the main student data in supervised school sessions. The same public link cannot stop forwarding or repeat submissions, so supervision is the practical control.",
        "Make the simple model the main model and the Transformer a second, experimental model. A small school dataset is more likely to support a simple model honestly.",
        "Split people, not trials. A student must appear only in training, validation, or final testing, never in more than one.",
        "Fix the app fields and consent process before main collection. The current app does not record role, grade, class, setting, or server-side completion.",
        "Treat reaction time as a secondary result unless the device and browser are kept similar for everyone.",
    ]:
        bullet(doc, text)

    heading(doc, "1. What was wrong with the earlier plan")
    heading(doc, "The link plan did not control who entered", 2)
    paragraph(doc, "The earlier plan said to choose individual students and privately send them one common link. That does not show that a response came from a chosen student. The link can be forwarded, and the current app creates a new anonymous session each time. Because the study will use only the normal link, the clean practical answer is to run it during supervised sessions with selected whole classes.")

    heading(doc, "The sampling plan was harder than necessary", 2)
    paragraph(doc, "Randomly choosing individual students across every grade and section would require many private messages, replacements, and permission checks. Choosing whole sections is easier for the school to schedule. It still gives each section a known chance of selection and follows a common method used in school surveys.")

    heading(doc, "The sample-size cutoffs were presented too confidently", 2)
    paragraph(doc, "The earlier numbers of 80, 150, and 300 were rough rules, not calculations for this model or target. A useful sample size depends on how common each answer is, how many model choices are tried, and how precise the final score must be. The corrected plan sets a practical class target and reports uncertainty instead of pretending that one number guarantees a valid model.")

    heading(doc, "The environment advice conflicted with the measurements", 2)
    paragraph(doc, "The earlier plan allowed both supervised sessions and unsupervised completion. That may be acceptable for choices, but it weakens reaction-time comparisons because devices, browsers, distractions, and input methods change timing. The corrected plan uses supervised student sessions and treats teacher or remote timing separately.")

    heading(doc, "The app does not collect the fields assumed by the plan", 2)
    paragraph(doc, "The current form asks for age group and highest education. It does not ask whether the person is a student or teacher, the student's grade or section, the testing setting, or the device. The 13-17 age group is too wide for the proposed student comparisons, and the education choices do not describe school grades well.")

    heading(doc, "The consent process is not strong enough for minors", 2)
    paragraph(doc, "The current app lets a minor tick a box saying that a guardian has consented. That records the student's statement; it does not prove that the school obtained parent or guardian consent. The school should collect the required permission before the session, and the app should then ask the student for their own agreement to take part.")

    heading(doc, "The current model evaluation is not a clean final test", 2)
    paragraph(doc, "The training code creates a participant-level training and validation split, which is good. But the data preparation is fitted before that split, so information such as response-time averages can come from validation participants. More seriously, the current held-out-task check loads an encoder trained on the full dataset and then tests a second model on some of the same participants. That score cannot be presented as performance on completely unseen people.")

    heading(doc, "The model claim was too broad", 2)
    paragraph(doc, "A model trained in one school during the weeks before exams cannot be called a general model of decision-making. Age, student or teacher role, exam stress, school culture, and testing conditions are mixed together. The corrected claim is limited to the selected classes and the collection period.")

    heading(doc, "2. The exact study question")
    paragraph(doc, "The recommended main question is: Using a student's results from the other four activities, can a simple model predict whether that student chooses the gamble in at least half of the risk activity trials?")
    paragraph(doc, "This is a clear yes-or-no target that matches the current project. Compare a simple model with the Transformer. The simple model should use a few summaries from the four activities. The Transformer may be reported as an experimental comparison, but it should not replace the simple model unless it clearly performs better on unseen students.")
    paragraph(doc, "The main unit is one student. The 240 trials give detail about that student; they do not turn one student into 240 independent participants.")

    heading(doc, "3. Who to collect data from")
    heading(doc, "Students", 2)
    numbered(doc, "List the classes", "Write down every eligible section in each participating grade and the number of students in each section.")
    numbered(doc, "Choose sections before collection", "Randomly choose one section from each grade. If the school can support more sessions, choose two. Also choose one backup section per grade before anyone sees the results.")
    numbered(doc, "Invite the whole selected section", "Give the information and permission form to every student in the selected section. Invite all students who return the required permission and agree to participate.")
    numbered(doc, "Use the normal link", "Share the ordinary app link in the selected class group shortly before the scheduled session or display it in the room. Do not share it in other class groups.")
    numbered(doc, "Record counts", "For each selected section, record how many students were enrolled, received the information, had permission, attended, started, completed, and stopped early. Keep this count sheet separate from the anonymous research data.")
    paragraph(doc, "A practical target is one full section from each participating grade. If four grades take part and a section has about 25 to 35 students, this gives roughly 100 to 140 invited students before absences and refusals. This is a target for planning, not a promise of model accuracy.")
    paragraph(doc, "Whole-class sampling is easier, but students from the same class may be more similar to one another than students from different classes. If the school allows it, include more sections rather than filling the sample from only one or two large sections. With only a few selected sections, do not claim that an overall percentage represents the whole school.")

    heading(doc, "Teachers", 2)
    paragraph(doc, "Invite all eligible teachers because there are few of them. Offer one or two staff-room sessions using the same link. Keep teacher data out of student model training. Report teacher completion and simple task summaries separately; only run a teacher model comparison if enough teachers complete it to avoid exposing individuals or producing meaningless percentages.")

    heading(doc, "4. Session plan")
    bullet(doc, "Run an 8 to 12 person pilot first. If instructions, task length, data fields, or code change afterward, do not mix pilot results into the main dataset.")
    bullet(doc, "Reserve a full 40 to 45 minute period until the pilot gives a reliable duration. The current app says roughly 30 minutes while the project configuration describes the full battery as roughly 40 minutes, so this must be made consistent.")
    bullet(doc, "Use the same room, device type, browser, input method, written instructions, and app version for student sessions as far as the school allows.")
    bullet(doc, "Keep the present task order if changing it safely would delay the study. State that tiredness may affect later tasks. Do not rush in a new order system immediately before collection.")
    bullet(doc, "Have a researcher read one standard introduction. Teachers should not stand behind students or see individual answers.")
    bullet(doc, "Offer one planned make-up session for consenting students who were absent. Do not replace absent students with convenient volunteers from other classes.")

    heading(doc, "5. Changes needed in the app before main collection")
    for text in [
        "Ask role first: Student or Teacher.",
        "For students, ask grade and a school-approved class label. For teachers, do not ask school grade or highest education unless it is needed for a stated question.",
        "For students, ask completed age or narrower choices. Report ages in wider groups later. For teachers, use broad age groups and include Prefer not to say.",
        "Replace the minor's guardian-confirmation checkbox with a message that school-held permission must already be recorded. Keep a separate student agreement checkbox in the app.",
        "Save completion status and end time to the database. At present completion is stored only in the browser, so exported data cannot cleanly show who finished.",
        "Save the collection group, supervised or unsupervised setting, and app version. Device and browser can be stored in a broad form without collecting a detailed fingerprint.",
        "Show One submission per person. The supervised attendance sheet is the control against repeats; do not use hidden browser fingerprinting.",
        "Keep names out of the research database. The school may keep a separate permission and attendance list, but it must not be joined to participant responses.",
    ]:
        bullet(doc, text)

    heading(doc, "6. Age groups and bias")
    paragraph(doc, "Use role first, then show the matching questions. For students, collect grade and completed age if approved. In the fair report, combine ages into 13-14, 15-16, and 17-18 only when each group is large enough to report safely. For teachers, broad ten-year age groups are enough, and teacher age may be left out when the group is very small.")
    paragraph(doc, "Do not claim that age caused a result. In this school, younger people will mostly be students and older people will mostly be teachers, so age and role cannot be separated. Report student results by grade or student age only. Treat teacher comparisons as descriptive.")

    heading(doc, "7. Training and testing without leakage")
    bullet(doc, "Freeze the question before training. Write down the target, model inputs, exclusions, score, and simple comparison model.")
    bullet(doc, "Split by participant. Place all trials from one person in one group. No participant may cross between training, tuning, and testing.")
    bullet(doc, "Fit preparation on training data. Response-time averages, scaling, vocabularies, missing-data rules, feature choices, and model settings must be learned inside the training data only.")
    bullet(doc, "Train the encoder only on training participants. The final test participants must not be used for masked training, early stopping, model choice, feature choice, or target choice.")
    bullet(doc, "Use the final test once. After every choice is fixed, evaluate once on the untouched students and save the result. Do not keep changing the model after seeing this score.")
    paragraph(doc, "If the dataset is too small to leave a useful final test group, use repeated five-part testing instead. Each round holds out different students, and every model choice must be made using only the other parts. Report the average score and how much it changes. Call this an internal check, not proof that the model works in other schools.")
    paragraph(doc, "The current evaluation code needs to be changed before its score is used in the fair. The encoder must be trained separately inside each outer split, and the data-preparation step must be fitted only on the training participants in that split.")

    heading(doc, "8. What to calculate")
    bullet(doc, "Recruitment: show enrolled, permitted, started, completed, and stopped counts for every selected section and for teachers.")
    bullet(doc, "Data quality: show completion time, missing trials, early stops, device or browser groups, and the number of repeated or suspicious sessions found during supervised review.")
    bullet(doc, "Behavior: show each task's main choice rate and typical response time. Use the middle response time rather than the average when a few very slow responses pull the average upward.")
    bullet(doc, "Prediction: show the simple model, the Transformer, and the rule that always predicts the more common answer.")
    bullet(doc, "Scores: show the confusion matrix, balanced accuracy, and a 95% range made by resampling whole participants. Explain balanced accuracy as giving equal importance to both answer groups.")
    bullet(doc, "Stability: repeat the participant-level split with several saved random seeds. If the score changes greatly, say that the model is unstable.")
    bullet(doc, "Class effect: show results for each selected section. If at least several sections are available, repeat the check while holding out whole sections. If only a few sections are available, limit the claim to those sections.")
    bullet(doc, "Subgroups: show grade or age results only when the groups are large enough. Do not create many small comparisons just to search for a positive result.")
    bullet(doc, "Partial sessions: report them. Do not silently delete everyone who completed fewer than 240 trials.")

    heading(doc, "9. Permission and privacy")
    paragraph(doc, "Get written school approval before collecting main data. For students under 18, use the parent or guardian permission and student agreement process required by the school or ethics authority. Participation must be voluntary, and refusing or stopping must not affect grades, attendance, or treatment by teachers.")
    paragraph(doc, "Explain the study purpose, approximate length, recorded data, possible discomfort or tiredness, right to stop, data access, storage period, and contact person in simple language. Keep the school permission list separate from anonymous app responses and decide in advance when both will be deleted.")

    heading(doc, "10. Timeline before the science fair")
    bullet(doc, "Day 1: ask for approval using this short plan. At the same time, list grades and sections and randomly choose the main and backup sections.")
    bullet(doc, "Days 1 to 3: make the minimum app changes, test the database export, and prepare parent information, student agreement, and a standard session script.")
    bullet(doc, "Days 3 and 4: run the pilot with 8 to 12 people. Freeze the app and analysis plan after fixing serious problems.")
    bullet(doc, "Days 5 to 12: run the selected class sessions and one make-up session. Run teacher sessions separately.")
    bullet(doc, "Days 13 and 14: close the dataset, save an untouched copy, check completion and duplicates, and create the participant-level splits.")
    bullet(doc, "Remaining days: run the fixed analysis, prepare the fair figures, and write limits beside every main result.")
    paragraph(doc, "If the school does not approve student collection in time, do not circulate the link informally to minors. Submit the working app, the pilot or teacher-only feasibility results that were properly approved, a tested analysis pipeline using demonstration data, and this future collection plan.")

    heading(doc, "11. Claims for the science fair")
    bullet(doc, "Allowed: We tested whether behavior in four activities predicted risk choices among participating students from selected classes in one school during the pre-exam period.")
    bullet(doc, "Allowed: The result is an internal school check and includes uncertainty across participant-level splits.")
    bullet(doc, "Allowed: Teacher results are a small separate description.")
    bullet(doc, "Not allowed: The model has been validated for all students, teachers, age groups, schools, or people.")
    bullet(doc, "Not allowed: The large number of trials replaces the need for many participants.")
    bullet(doc, "Not allowed: A held-out activity proves that the model understands a completely new activity. The training process has already used the activity type unless it is removed from encoder training as well.")

    heading(doc, "12. Launch checklist")
    for text in [
        "The school has approved the study, parent permission process, schedule, and data fields.",
        "Main and backup class sections were chosen randomly before collection.",
        "The link will be used only in selected class sessions and teacher sessions.",
        "The app records role, student grade or teacher status, revised age, setting, app version, and server-side completion.",
        "The app's stated duration matches the pilot result.",
        "The final research question, outcome, simple model, Transformer settings, exclusions, and scores are written down.",
        "The code keeps every participant entirely inside one split and fits all preparation inside training data.",
        "The report will include recruitment counts, partial sessions, uncertainty, and limits.",
    ]:
        bullet(doc, text)

    heading(doc, "Sources")
    sources = [
        ("ICMR: ethical guidance for research involving children", "https://www.icmr.gov.in/icmrobject/custom_data/pdf/resource-guidelines/National_Ethical_Guidelines_for_BioMedical_Research_Involving_Children_0.pdf"),
        ("CDC: school surveys that select whole classes", "https://www.cdc.gov/mmwr/preview/mmwrhtml/rr5312a1.htm"),
        ("WHO school survey example using selected classes", "https://cdn.who.int/media/docs/default-source/ncds/ncd-surveillance/data-reporting/sri-lanka/gshs_sri_lanka_2024_report.pdf"),
        ("AAPOR: selection and inference with non-probability samples", "https://aapor.org/wp-content/uploads/2023/02/Task-Force-Report-FINAL.pdf"),
        ("TRIPOD+AI: reporting model development and evaluation", "https://www.bmj.com/content/385/bmj-2023-078378"),
        ("Scikit-learn: avoiding leakage and using nested testing", "https://scikit-learn.org/stable/auto_examples/model_selection/plot_nested_cross_validation_iris.html"),
        ("Bridges and colleagues: browser and device timing differences", "https://pmc.ncbi.nlm.nih.gov/articles/PMC7512138/"),
        ("Riley and colleagues: sample size depends on the prediction problem", "https://www.bmj.com/content/368/bmj.m441"),
    ]
    for label, url in sources:
        p = doc.add_paragraph()
        add_hyperlink(p, label, url)

    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
        if p.style.name.startswith("Heading"):
            p.paragraph_format.keep_with_next = True

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
