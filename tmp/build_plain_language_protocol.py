from datetime import date
from pathlib import Path

from docx import Document

from build_simple_protocol import (
    add_hyperlink,
    bullet,
    configure_styles,
    heading,
    numbered,
    paragraph,
    set_font,
)


ROOT = Path(r"C:\Users\USER\Desktop\code\Neuroscope")
OUT = ROOT / "output" / "documents" / "neuroscope-data-collection-plan-plain-language.docx"


def add_title(doc):
    p = doc.add_paragraph()
    r = p.add_run("NEUROSCOPE Data Collection Plan")
    set_font(r, size=20, bold=True)

    p = doc.add_paragraph()
    r = p.add_run("A practical plan for collecting student and teacher data before the science fair")
    set_font(r, italic=True)

    p = doc.add_paragraph()
    r = p.add_run("Prepared: ")
    set_font(r, bold=True)
    r = p.add_run(date.today().strftime("%d %B %Y"))
    set_font(r)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "NEUROSCOPE Data Collection Plan"
    doc.core_properties.author = "NEUROSCOPE research team"
    add_title(doc)

    heading(doc, "Short answer")
    for text in [
        "Do not post the app link openly in every class group. Students who are more interested, more confident, or more available will be more likely to take part, so the results may not represent the school.",
        "Choose students from every grade and section first. Then send the normal app link directly to the chosen students.",
        "Send the same link privately to every eligible teacher. Keep teacher results separate because the teacher group will be much smaller.",
        "Run a small test with 8 to 12 people immediately. Fix confusing instructions and technical problems, then begin the main collection before the science-fair deadline.",
        "Use several time slots instead of depending only on zero period. Keep the instructions and testing conditions as similar as possible.",
        "Divide the data by person. All answers from one person must stay together in either the training group or the testing group.",
        "If the final student sample is small, present the project as an early study or a working prototype. Do not claim that the model has been proven for all students, teachers, or people in general.",
    ]:
        bullet(doc, text)

    heading(doc, "1. The decision")
    paragraph(doc, "The main study should be about students in this school. The model should be trained and tested mainly on student data. Teachers should be treated as a separate small group that shows how the app behaves with adults, but a small teacher group cannot prove that the model works for adults in general.")
    paragraph(doc, "Having more students than teachers is not automatically unfair or incorrect. A school really does have more students. The problem begins when student and teacher data are mixed into one result, because the large student group will control that result. Report the two groups separately.")

    heading(doc, "2. How to choose people")
    heading(doc, "Students", 2)
    numbered(doc, "Make a list", "Ask the school for the number of eligible students in each grade and section. The school can keep the names; the research file does not need them.")
    numbered(doc, "Set a target", "Decide how many students you can realistically test. Divide that target across grades and sections so that every important group is included. A larger class can receive more places than a smaller class.")
    numbered(doc, "Choose at random", "Use shuffled roll numbers or another random method inside each grade and section. Also choose a short replacement list from the same group for students who are absent or do not agree to participate.")
    numbered(doc, "Send the link directly", "Send the ordinary app link only to the students who were chosen. This can be done through a private message, a teacher, or the class representative. A class-group message may remind the chosen students, but it should clearly say that only selected students should participate.")
    numbered(doc, "Count each stage", "For every grade or section, record how many students were eligible, chosen, allowed by a parent or guardian, started, completed, and stopped. These counts show whether some groups were missed.")

    heading(doc, "Teachers", 2)
    paragraph(doc, "Invite every eligible teacher because the teacher group is already small. Send the same app link privately and offer several possible times. Report how many teachers were eligible, invited, started, and completed the study. Do not create fake teacher data or copy teacher records to make the groups look equal.")

    heading(doc, "3. Where and when to collect the data")
    paragraph(doc, "The best realistic option is a mixed plan. Run a few short supervised sessions in a quiet room during approved zero periods, free periods, lunch, or other low-disruption times. Chosen people who cannot attend may use the same link elsewhere during the collection window, but record where and on what device they completed it.")
    paragraph(doc, "This will not be a perfect laboratory experiment, and that is acceptable if it is described honestly. Try to keep the written instructions, practice questions, app version, browser, device type, and task order the same. Record the date, time, place, device, interruptions, and whether the person was supervised.")

    heading(doc, "Suggested schedule", 2)
    for text in [
        "Days 1 and 2: get school approval, finish the student-selection plan, and freeze the main app questions.",
        "Days 2 and 3: test the process with 8 to 12 people. Check how long it takes, which instructions are unclear, where people quit, and whether the data saves correctly.",
        "Days 3 to 12: collect the main data across several approved time slots.",
        "Final 2 or 3 days: stop collection at the announced time, keep an untouched copy of the data, run the planned analysis, and write the report.",
    ]:
        bullet(doc, text)

    heading(doc, "4. Age questions in the app")
    paragraph(doc, "Ask whether the person is a student or a teacher before asking about age. Show different questions to the two groups.")
    bullet(doc, "For students, record grade and completed age if the school approves it. In the report, use groups such as 13-14, 15-16, and 17-18 so that individual students are harder to identify.")
    bullet(doc, "For teachers, use broad groups such as 20-29, 30-39, 40-49, 50-59, and 60+, or leave teacher age out if very few teachers participate.")
    bullet(doc, "Do not use age to guess whether someone is a student or teacher. Store the role directly.")
    paragraph(doc, "Age and role will overlap very little in one school, so the study cannot tell whether a student-teacher difference is caused by age, experience, or the role itself. State this as a limit instead of claiming that the model has discovered an age effect.")

    heading(doc, "5. Training and testing the model")
    paragraph(doc, "Your idea of keeping some students aside for testing is good, but the testing students must remain completely unused until the model choices are finished. A very small testing group can give a lucky or unlucky score, so it does not prove that the result is correct by itself.")
    for text in [
        "Split by person, not by answer or trial. If one person completes 240 trials, all 240 must stay in the same group.",
        "Use the training group to choose settings and improve the model. Use the final testing group only once, after those choices are fixed.",
        "Keep student and teacher results separate. Testing on a few teachers is an early check, not proof that the model works for teachers everywhere.",
        "Compare the model with a simple rule, such as always predicting the most common answer. The model should clearly beat that simple rule.",
        "Report a range around each score, because one exact percentage makes a small study look more certain than it is.",
    ]:
        bullet(doc, text)

    heading(doc, "What to do at different sample sizes", 2)
    bullet(doc, "Fewer than 80 completed students: treat the project as an early test of whether the app works. Show that useful patterns may exist, but avoid strong accuracy claims.")
    bullet(doc, "80 to 149 completed students: repeatedly train on most students and test on the remaining students, changing the groups each time. Report the average result and how much it changes.")
    bullet(doc, "150 to 299 completed students: keep about 15% of the students untouched for the final test, as long as all important grades are still represented.")
    bullet(doc, "300 or more completed students: keep about 20% untouched for the final test. A later test in another school would still be needed before making wider claims.")
    paragraph(doc, "These numbers are practical guides, not guarantees. The right number also depends on how balanced the answers are, how difficult the prediction is, and how complex the model is.")

    heading(doc, "6. Biases to watch")
    bullet(doc, "Volunteer bias: an open link mainly reaches people who choose themselves. Selecting students first reduces this problem.")
    bullet(doc, "Missing people: absent students and people who do not finish may be different from those who complete the study. Show selection and completion numbers for each grade or section.")
    bullet(doc, "Exam timing: stress and tiredness may affect answers. Record the date and time, and say clearly that the study happened shortly before exams.")
    bullet(doc, "Different settings: room noise, phones, computers, browsers, and interruptions can change the results. Keep them similar when possible and record the differences.")
    bullet(doc, "Task order and tiredness: later tasks may receive worse answers because people are tired. Use different task orders if the app supports this safely; otherwise use one order and mention this limit.")
    bullet(doc, "Dropouts: keeping only perfect completed sessions hides where people stopped. Save the last completed task and the reason for any known technical failure.")

    heading(doc, "7. Simple analysis methods")
    bullet(doc, "Show the number invited, started, and completed in every grade or section. This is the clearest check for missing groups.")
    bullet(doc, "If one class had a much higher chance of being chosen than another, give the underrepresented class more weight when describing the whole student population. Also show the result without this adjustment.")
    bullet(doc, "Build score ranges by repeatedly drawing whole participants from the data and recalculating the result. Never draw individual trials as if they came from different people.")
    bullet(doc, "When studying trial-by-trial answers, use a method that knows that many answers came from the same person. Two hundred and forty answers from one student are still data from one student, not 240 students.")
    bullet(doc, "Use measures that work when one answer is more common than another, such as balanced accuracy or F1, and explain them in one sentence in the fair report.")
    bullet(doc, "Run the analysis once with completed sessions only and again while including information about where incomplete sessions stopped. If the conclusion changes, report that openly.")

    heading(doc, "8. Permission and privacy")
    paragraph(doc, "The short deadline does not remove the need for school approval. For students under 18, follow the parent or guardian consent and student agreement process required by the school or ethics authority. A checkbox inside the app is not enough unless the school has approved that process.")
    paragraph(doc, "Participation must not affect grades, attendance, teacher evaluation, or access to school activities. Keep names outside the research dataset. Do not publish very small groups that could reveal a person, and decide who may open the data and when it will be deleted.")

    heading(doc, "9. What the final report can say")
    bullet(doc, "Safe claim: the model found and tested patterns among the participating students from this school during the pre-exam period.")
    bullet(doc, "Safe claim: teacher results were a small separate check and should be treated as early evidence.")
    bullet(doc, "Unsafe claim: the model is proven for all students, all teachers, every age group, or the general public.")
    bullet(doc, "Unsafe claim: a small held-back group removes all bias. It only checks performance on those held-back people.")

    heading(doc, "10. Checklist before starting")
    for text in [
        "The school has approved the study, the schedule, and the consent process.",
        "Students have been selected across grades and sections before the link is sent.",
        "The normal link will be sent privately to selected students and all eligible teachers.",
        "Replacement students come from the same grade or section as the person they replace.",
        "The app records role, grade, age group, date, time, setting, device, app version, and last completed task.",
        "All data from one person stays in one training or testing group.",
        "The stopping date, excluded data, scores, simple comparison rule, and allowed claims are written down before the final analysis.",
    ]:
        bullet(doc, text)

    heading(doc, "Sources")
    sources = [
        ("ICMR guidance for research involving children", "https://www.icmr.gov.in/icmrobject/custom_data/pdf/resource-guidelines/National_Ethical_Guidelines_for_BioMedical_Research_Involving_Children_0.pdf"),
        ("STROBE guidance for reporting observational studies", "https://www.equator-network.org/reporting-guidelines/strobe/"),
        ("AAPOR report on samples where people choose whether to join", "https://aapor.org/wp-content/uploads/2022/11/NPS_TF_Report_Final_7_revised_FNL_6_22_13.pdf"),
        ("Riley and colleagues on sample size for prediction models", "https://www.bmj.com/content/368/bmj.m441"),
        ("Varoquaux on small datasets and changing test scores", "https://pubmed.ncbi.nlm.nih.gov/28655633/"),
        ("TRIPOD+AI guidance for reporting prediction models", "https://pmc.ncbi.nlm.nih.gov/articles/PMC11019967/"),
        ("Bridges and colleagues on timing in online and laboratory tasks", "https://pubmed.ncbi.nlm.nih.gov/33005482/"),
        ("Sievertsen and colleagues on tiredness and student performance", "https://pmc.ncbi.nlm.nih.gov/articles/PMC4790980/"),
    ]
    for label, url in sources:
        p = doc.add_paragraph(style="List Bullet")
        add_hyperlink(p, label, url)

    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
        if p.style.name.startswith("Heading"):
            p.paragraph_format.keep_with_next = True

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
