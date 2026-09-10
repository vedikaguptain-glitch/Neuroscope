from datetime import date
from pathlib import Path

from docx import Document

from build_simple_protocol import (
    add_hyperlink,
    bullet,
    configure_styles,
    heading,
    numbered,
    paragraph,
    set_font,
)


ROOT = Path(r"C:\Users\USER\Desktop\code\Neuroscope")
OUT = ROOT / "output" / "documents" / "neuroscope-student-thinking-study-plan.docx"


def add_title(doc):
    p = doc.add_paragraph()
    r = p.add_run("NEUROSCOPE Student Thinking Study Plan")
    set_font(r, size=20, bold=True)

    p = doc.add_paragraph()
    r = p.add_run("Students aged 15 to 18 in grades 9 to 12")
    set_font(r, italic=True)

    p = doc.add_paragraph()
    r = p.add_run("Prepared: ")
    set_font(r, bold=True)
    r = p.add_run(date.today().strftime("%d %B %Y"))
    set_font(r)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "NEUROSCOPE Student Thinking Study Plan"
    doc.core_properties.author = "NEUROSCOPE research team"
    add_title(doc)

    heading(doc, "Final decision")
    paragraph(doc, "The main study will focus only on student decision patterns. Eligible students will be aged 15, 16, 17, or 18 and enrolled in grades 9, 10, 11, or 12. Teachers will not be part of the main sample or main model.")
    paragraph(doc, "The student sample will come from three sections in each grade. From each section, choose 10 interested and eligible students. This gives 10 students x 3 sections x 4 grades = 120 planned student participants.")
    for text in [
        "Target: 120 completed student sessions.",
        "Minimum for the main analysis: 80 completed student sessions.",
        "Choose interested volunteers, then randomly select among them when more than 10 volunteer in a section.",
        "Use the normal app link only during the planned sessions.",
        "Run groups of about 10 students in zero period or another approved low-disruption time.",
        "Treat the teacher study as a separate optional extension after the student study.",
        "The only major blocker is school permission. Do not collect data from minors without it.",
    ]:
        bullet(doc, text)

    heading(doc, "1. What the study is about")
    paragraph(doc, "The main question is: What decision patterns appear among interested students aged 15 to 18, and how do those patterns change across the four ages?")
    paragraph(doc, "The five activities can show different parts of decision-making:")
    bullet(doc, "Probabilistic learning: how quickly students learn which choice is more rewarding and how they adjust when the pattern changes.")
    bullet(doc, "Risk preference: how often students choose a gamble instead of a safe option.")
    bullet(doc, "Delay discounting: how often students choose a larger later reward instead of a smaller immediate reward.")
    bullet(doc, "Rule discovery: how accurately students find and switch between rules.")
    bullet(doc, "Social ultimatum: how students respond to fair and unfair offers.")
    paragraph(doc, "Reaction time can be reported as supporting information, but the main results should be the choices and accuracy because device and browser differences can change timing.")
    paragraph(doc, "This study is about decision patterns, not intelligence, mental health, personality, or academic ability. Do not describe a participant or age group as smarter, better, or worse.")

    heading(doc, "2. Who is included")
    bullet(doc, "The participant is a current student in grade 9, 10, 11, or 12.")
    bullet(doc, "The participant is exactly 15, 16, 17, or 18 years old on the day of the session.")
    bullet(doc, "The participant is interested in taking part.")
    bullet(doc, "The school has approved participation and the required parent or guardian permission has been completed for minors.")
    bullet(doc, "The student personally agrees to participate and may stop at any time.")
    paragraph(doc, "Students outside ages 15 to 18 are not part of the main study even if they are in grades 9 to 12. Record how many volunteers were ineligible because of age, but do not collect their task data for the main dataset.")

    heading(doc, "3. How to select the 120 students")
    numbered(doc, "Choose sections", "Select three sections from each of grades 9, 10, 11, and 12. If a grade has more than three sections, choose the three sections randomly before recruitment.")
    numbered(doc, "Ask for interest", "Give the same short announcement to all students in each selected section. Ask eligible students who are interested to put their names on a school-held volunteer list. Do not send the experiment link at this stage.")
    numbered(doc, "Check eligibility", "Keep only volunteers aged 15 to 18 who are in the correct grade and can complete the required permission process.")
    numbered(doc, "Select 10", "If more than 10 eligible students volunteer in a section, choose 10 by a random draw. Do not select based on marks, perceived intelligence, friendship, behavior, or who a teacher prefers.")
    numbered(doc, "Choose backups", "Randomly choose two or three backup volunteers from the same section. Use them only when a selected student cannot participate or does not return permission.")
    numbered(doc, "Repeat for every section", "Ten students from three sections in each of four grades gives 120 main selections. The backup list protects the target when someone is absent or stops before starting.")
    paragraph(doc, "Interest is an intentional part of this design. The result will describe interested students from the selected sections. It must not be presented as a perfectly representative result for every student in the school.")
    paragraph(doc, "If a selected section has fewer than 10 eligible volunteers, first repeat the same announcement once. If the number is still low, use a section selected in advance as a backup from the same grade. Do not quietly fill the places with friends or high-performing students.")

    heading(doc, "4. Age and grade balance")
    paragraph(doc, "The app should ask exact age with four main choices: 15, 16, 17, and 18. It should separately ask grade: 9, 10, 11, or 12. Age and grade are related but not identical, so both should be stored.")
    paragraph(doc, "Try to reach about 25 to 35 completed students at each age. Do not force equal age groups by rejecting eligible students after seeing their answers. Balance the ages using the volunteer list before the experiment begins.")
    paragraph(doc, "Show results for age 15, 16, 17, and 18 separately. Also show the number of students in each age and grade. With roughly 30 students per age, use one overall age-trend test instead of running many pair-by-pair tests that can produce chance findings.")
    paragraph(doc, "Do not train four separate Transformers, one for each age. Each age group will be too small for that. Train the student model on the full student sample and use age only for planned comparisons and error checks.")

    heading(doc, "5. Session plan")
    paragraph(doc, "Run about 10 students at a time. With 12 selected sections, this means about 12 student sessions unless the school allows two groups to run at the same time. A session can be held in zero period, a free period, an activity period, or another approved slot.")
    for text in [
        "Run an 8 to 12 person pilot before the main sessions. If the app or instructions change afterward, keep pilot data out of the main dataset.",
        "Reserve 40 to 45 minutes until the pilot gives a reliable duration. Make the duration shown in the app match the pilot result.",
        "Use the same room, device type, browser, input method, written instructions, task order, and app version as far as possible.",
        "Share the normal app link at the start of the supervised session. Do not float it across general class groups.",
        "Keep an attendance and permission list at the school, but do not connect names on that list to anonymous app responses.",
        "Teachers should not stand behind students or see individual answers.",
        "Offer one or two make-up sessions for selected students who had permission but were absent.",
    ]:
        bullet(doc, text)

    heading(doc, "6. How to ask the principal")
    paragraph(doc, "Ask for approval for the full plan once rather than requesting separate permission every morning. The request should say that only about 10 students from a selected section will leave normal activities for one supervised 40 to 45 minute session.")
    paragraph(doc, "The approval request should include:")
    bullet(doc, "Grades 9 to 12 only, ages 15 to 18.")
    bullet(doc, "About 10 students per selected section and 120 students in total.")
    bullet(doc, "Zero period or another time chosen by the school to avoid lesson disruption.")
    bullet(doc, "Voluntary participation, required parent or guardian permission, and student agreement.")
    bullet(doc, "No names stored with experiment responses and no effect on marks or attendance.")
    bullet(doc, "A short pilot, a fixed session script, and the right to stop participation.")
    paragraph(doc, "Exam timing remains a limitation because stress and tiredness may affect choices. Record the date and session time and describe the study as a pre-exam study. If the principal does not approve the student sessions, do not collect student data informally; the student study cannot proceed ethically without permission.")

    heading(doc, "7. Information the app must collect")
    for text in [
        "Role: Student for the main study. Keep Teacher only for the separate extension.",
        "Exact age choice: 15, 16, 17, or 18 for the student study.",
        "Grade: 9, 10, 11, or 12.",
        "A school-approved section label that does not contain a student's name.",
        "Session group, supervised setting, app version, start time, end time, and completion status.",
        "A clear One submission per person message.",
        "No participant name in the research database.",
    ]:
        bullet(doc, text)

    heading(doc, "8. Permission and privacy")
    paragraph(doc, "The school should collect the required parent or guardian permission before a minor attends the session. The app should ask the student for their own agreement to take part. A student should not be asked to prove guardian permission by ticking a box inside the app.")
    paragraph(doc, "Participation must be voluntary. Refusing or stopping must not affect marks, attendance, teacher treatment, or access to school activities. Explain the purpose, length, recorded data, possible tiredness, right to stop, data access, storage period, and contact person in simple language.")

    heading(doc, "9. Analysis")
    heading(doc, "Student thinking by age", 2)
    bullet(doc, "For each age, report the number of students and the main result from each of the five activities.")
    bullet(doc, "Use percentages for choices and accuracy. Use the middle response time rather than the average when a few very slow responses pull the average upward.")
    bullet(doc, "Show a 95% uncertainty range around each age result. Treat the participant, not each trial, as the independent unit.")
    bullet(doc, "Test one planned age trend from 15 through 18 for each main task result. Clearly label any extra comparisons as exploratory.")
    bullet(doc, "Compare age groups only after checking that each group contains enough students. If an age group is very small, show it descriptively without a strong conclusion.")

    heading(doc, "Model testing", 2)
    paragraph(doc, "Use repeated five-part testing for the student model. Divide participants into five groups, train on four groups, test on the remaining unseen group, and rotate until every group has been tested. Repeat with several saved random splits and report the average result and how much it changes.")
    paragraph(doc, "All trials from one student must stay together. Data preparation, response-time scaling, feature selection, encoder training, early stopping, and model settings must use training students only. The simple model is the main comparison; the Transformer is secondary.")

    heading(doc, "10. Teacher extension")
    paragraph(doc, "The teacher extension is a separate test study. Invite interested teachers only after the student collection plan is secure. Teacher data must not be added to the 120-student sample, used to balance student ages, or mixed into the student model's main score.")
    paragraph(doc, "Report teacher participation and simple task summaries. If the teacher group is small, do not report detailed age groups or train a separate teacher model. The extension can show whether the procedure is workable with adults, but it cannot prove that student and teacher thinking differs.")

    heading(doc, "11. Timeline")
    bullet(doc, "Day 1: submit the full permission request and list the available sections in grades 9 to 12.")
    bullet(doc, "Days 1 to 3: make the app changes, prepare the volunteer message and permission materials, and test the data export.")
    bullet(doc, "Days 3 and 4: run the pilot and then freeze the app, instructions, and analysis plan.")
    bullet(doc, "Days 5 to 8: collect volunteer names, check ages, complete permission, and randomly choose 10 plus backups from each section.")
    bullet(doc, "Days 6 to 17: run about 12 student sessions and the make-up sessions. Run parallel sessions only if supervision and devices are adequate.")
    bullet(doc, "Final days: close the data, save an untouched copy, check completion and repeats, run the fixed analysis, and prepare the science-fair report.")

    heading(doc, "12. Claims for the science fair")
    bullet(doc, "Allowed: We studied decision patterns among interested students aged 15 to 18 from selected sections in grades 9 to 12 at one school.")
    bullet(doc, "Allowed: We compared planned task results across exact ages and tested the model on unseen students.")
    bullet(doc, "Allowed: The data were collected shortly before exams, which may have affected the results.")
    bullet(doc, "Not allowed: The sample represents every student in the school.")
    bullet(doc, "Not allowed: The tasks measure intelligence, personality, mental health, or academic potential.")
    bullet(doc, "Not allowed: A small teacher extension proves a student-teacher difference.")

    heading(doc, "13. Final checklist")
    for text in [
        "The main study contains only students aged 15, 16, 17, or 18 in grades 9 to 12.",
        "Three sections per grade have been selected, giving 12 section groups.",
        "Ten interested students plus two or three backups are randomly chosen from each section.",
        "The target is 120 completed student sessions and the minimum is 80.",
        "Exact age, grade, section label, setting, app version, and completion are recorded.",
        "The principal has approved the full plan and the required permission process is complete.",
        "The normal link is shared only at the supervised sessions.",
        "Age comparisons are planned before data analysis.",
        "Every student's trials remain inside one training or testing group.",
        "Teacher data remain a separate optional extension.",
    ]:
        bullet(doc, text)

    heading(doc, "Sources")
    sources = [
        ("ICMR: ethical guidance for research involving children", "https://www.icmr.gov.in/icmrobject/custom_data/pdf/resource-guidelines/National_Ethical_Guidelines_for_BioMedical_Research_Involving_Children_0.pdf"),
        ("CDC: school surveys that select classes and obtain permission", "https://www.cdc.gov/mmwr/preview/mmwrhtml/rr5312a1.htm"),
        ("AAPOR: limits of volunteer and other non-probability samples", "https://aapor.org/wp-content/uploads/2023/02/Task-Force-Report-FINAL.pdf"),
        ("TRIPOD+AI: separating model development and testing data", "https://www.bmj.com/content/385/bmj-2023-078378"),
        ("Bridges and colleagues: browser and device timing differences", "https://pmc.ncbi.nlm.nih.gov/articles/PMC7512138/"),
    ]
    for label, url in sources:
        p = doc.add_paragraph()
        add_hyperlink(p, label, url)

    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
        if p.style.name.startswith("Heading"):
            p.paragraph_format.keep_with_next = True

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
