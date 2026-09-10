from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\USER\Desktop\code\Neuroscope")
OUT = ROOT / "output" / "documents" / "neuroscope-data-collection-plan-simple.docx"
BLACK = "000000"


def set_font(run, size=11, bold=None, italic=None):
    run.font.name = "Arial"
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_font(run, size=9)


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)
    doc.settings.odd_and_even_pages_header_footer = False

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for level, size, before, after in ((1, 14, 12, 5), (2, 12, 9, 4)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.05

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)


def add_title(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("NEUROSCOPE: Data Collection and Bias-Reduction Plan")
    set_font(r, size=20, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("School-based study of students and teachers")
    set_font(r, size=11, italic=True)

    p = doc.add_paragraph()
    r = p.add_run("Prepared: ")
    set_font(r, bold=True)
    r = p.add_run(date.today().strftime("%d %B %Y"))
    set_font(r)


def heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def paragraph(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_font(r)
    return p


def numbered(doc, lead, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(lead + ": ")
    set_font(r, bold=True)
    r = p.add_run(text)
    set_font(r)
    return p


def add_hyperlink(paragraph, label, url):
    rid = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    rpr.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLACK)
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "NEUROSCOPE Data Collection and Bias-Reduction Plan"
    doc.core_properties.author = "NEUROSCOPE research team"

    add_title(doc)

    heading(doc, "TL;DR")
    for text in [
        "Do not use an open class-group link as the main sampling method. It will mainly attract students who are interested, available, and comfortable with the task.",
        "Use the school roster to randomly select students within each grade and section. Send the app only to selected students using invitation codes.",
        "Invite all eligible teachers, but analyze teachers separately from students. A small teacher group cannot support strong conclusions about adults.",
        "Run a short pilot immediately, then collect the main dataset before the science-fair cutoff. Use several permitted time blocks rather than only zero period.",
        "Use the same room, device type, browser, instructions, and practice procedure for everyone as far as possible.",
        "Ask whether the participant is a student or teacher before asking age. Use narrow student age groups and broad teacher age groups.",
        "Keep every trial from one participant in the same training or evaluation split. Never split one person's trials across both.",
        "If fewer than about 150 students complete the study, describe the model as preliminary or a feasibility result, not a validated model of human decision-making.",
        "For participants under 18, use the guardian-consent and student-assent process required by the school or ethics authority.",
    ]:
        bullet(doc, text)

    heading(doc, "1. Main decision")
    paragraph(doc, "The primary study population should be students enrolled in the school. The main result should describe how well NEUROSCOPE learns patterns among new students from this school. It should not be described as a general model of all humans.")
    paragraph(doc, "Teachers should form a separate exploratory group. Train and evaluate the primary model on students. Then test the student-developed pipeline on teachers without using the teacher records during primary training. If the teacher sample is small, report the individual results or summary cautiously with confidence intervals.")
    paragraph(doc, "The fact that there are more students than teachers is not itself a bias; it is the real population structure of a school. The problem occurs when one combined score is presented as if the two groups were equally represented or interchangeable.")

    heading(doc, "2. Recruitment and data collection")
    heading(doc, "Student sampling", 2)
    for lead, text in [
        ("Create a sampling list", "Obtain the number of eligible students in each grade and section. The school can keep names; the research dataset only needs anonymous invitation IDs."),
        ("Select within groups", "Randomly select students inside each grade and section. Draw an ordered list of replacements from the same group in case someone is absent or does not consent."),
        ("Use private invitations", "Give selected participants one-time invitation codes. A class-group message may remind selected students, but it should not allow unsampled students to enter the main study."),
        ("Record recruitment counts", "For each grade, record eligible, invited, guardian-approved, assented or consented, started, completed, and excluded counts."),
    ]:
        numbered(doc, lead, text)

    heading(doc, "Teacher recruitment", 2)
    paragraph(doc, "Invite all eligible teachers privately and offer several time slots. The principal should not receive a list of teachers who participated. Report the number eligible, invited, and completed. Do not copy teacher sessions, generate synthetic teacher records, or split a teacher's trials across folds to make the sample appear balanced.")

    heading(doc, "Collection schedule before the science fair", 2)
    for text in [
        "Days 1-2: obtain approval, freeze the sampling plan, and make the minimum app changes.",
        "Days 2-3: run an 8-12 person pilot to measure actual duration, misunderstood instructions, technical failures, and dropout points. Pilot results are for fixing the procedure, not proving model accuracy.",
        "Days 3-12: run the main sessions across permitted zero periods, free periods, lunch, or other low-disruption blocks. Zero period should not be the only option.",
        "Two or three days before submission: stop collection at a declared time, freeze the raw data and code version, and perform the predeclared analysis.",
        "If the final sample is small, lead the science-fair report with the working system, data-quality results, descriptive patterns, and uncertainty rather than one optimistic accuracy score.",
    ]:
        bullet(doc, text)

    heading(doc, "Standard session conditions", 2)
    for text in [
        "Use a quiet room, the same device class and browser, the same input method, and the same written instructions.",
        "Provide the same practice trials and do not include practice trials in the analysis.",
        "Record the date, time block, room, device type, browser, interruptions, app version, and task order.",
        "Protect the participant's screen and avoid teachers watching student choices, especially during the social task.",
        "If possible, assign the five tasks in balanced orders. If that change cannot be tested safely before collection, keep one fixed order for everyone and state that task and position effects cannot be separated.",
    ]:
        bullet(doc, text)

    heading(doc, "3. Age and demographic fields")
    paragraph(doc, "The current 13-17 category is too wide for a school study because decision behavior can change during adolescence. The current adult education question also provides little useful variation among students.")
    paragraph(doc, "Ask role first: Student or Teacher. Then show different questions for each group.")
    for text in [
        "Students: record grade or year, a pseudonymous section code, and completed age if approved. Ages can be reported as 13-14, 15-16, and 17-18 to protect privacy.",
        "Teachers: use broad age groups such as 20-29, 30-39, 40-49, 50-59, and 60+, or omit age if the staff is so small that people could be identified.",
        "Store role directly. Do not infer student or teacher status from age or education.",
        "Do not publish small teacher or student subgroups that could identify a person. Combine or suppress very small cells.",
    ]:
        bullet(doc, text)

    heading(doc, "4. Training and validation")
    paragraph(doc, "A few students held aside after training cannot prove that the model works for teachers. Validation must match the population being claimed, and the evaluation participants must remain unused during model fitting and tuning.")
    for text in [
        "Split by participant, never by trial. All 240 trials from one person must stay in the same outer fold or test set.",
        "Fit normalization, vocabularies, feature selection, and model tuning only on the training portion inside each fold.",
        "Use validation folds for hyperparameters and early stopping. Use a locked test set only after the model family and analysis plan are fixed.",
        "Report the majority baseline, balanced accuracy, macro-F1 where relevant, and participant-level confidence intervals. Do not report only ordinary accuracy.",
        "Repeat the evaluation across several participant-level splits or use nested cross-validation when the sample is modest.",
        "The current held-out-task probe is cross-task prediction, because the encoder was pretrained on all five tasks. It is not evidence that the encoder can handle a completely unseen task.",
    ]:
        bullet(doc, text)

    heading(doc, "Practical interpretation by completed student sample", 2)
    for text in [
        "Below 80: feasibility study only. Use simple baselines and repeated participant-level evaluation; expect very wide uncertainty.",
        "80-149: exploratory model comparison using repeated nested cross-validation. Do not call the Transformer validated.",
        "150-299: if all important grades remain represented, lock about 15% as an untouched test set and tune only on the remaining data.",
        "300 or more: lock about 20% as a test set and seek later evaluation in another school or adult cohort.",
    ]:
        bullet(doc, text)
    paragraph(doc, "These ranges are decision rules, not formal power calculations. Final sample size depends on the target, label balance, model complexity, expected performance, and required precision.")

    heading(doc, "5. Bias controls and statistical methods")
    heading(doc, "Main biases to report", 2)
    for text in [
        "Self-selection: controlled mainly through roster-based random invitations instead of an open link.",
        "Nonresponse: measured by comparing invitation and completion rates across grades and sections.",
        "Teacher nonresponse: reported against the total number of eligible teachers.",
        "Exam-period effects: recorded through the date and session block; conclusions must be limited to the pre-exam setting.",
        "Order and fatigue: reduced through balanced task order where feasible, or disclosed as a limitation when the order is fixed.",
        "Device and environment: reduced through standard devices and recorded as metadata.",
        "Attrition: measured by retaining the last completed task and reason for technical failure. The current extraction of only perfect 240-trial sessions hides this problem unless partial-session metadata is kept separately.",
        "Age-role confounding: handled by analyzing students and teachers separately. The school sample cannot separate age effects from role effects because the groups barely overlap in age.",
    ]:
        bullet(doc, text)

    heading(doc, "Useful statistical techniques", 2)
    for text in [
        "Use sampling weights for student population summaries when grades or sections were sampled at different rates. A base weight is the inverse of the probability of selection.",
        "Adjust weights for different response rates only within predeclared groups, and report weighted and unweighted results. Weighting cannot fix differences that were never measured.",
        "Report effective sample size after weighting: n_eff = (sum of weights)^2 / sum of squared weights.",
        "If whole classes are sampled, account for within-class similarity. A common planning approximation is design effect = 1 + (average class size - 1) x ICC.",
        "Use mixed-effects or hierarchical models for trial-level outcomes because 240 trials from one participant are repeated measurements, not 240 independent people.",
        "Use participant-level bootstrap confidence intervals and permutation tests that preserve participant grouping.",
        "Compare complete-session analysis with a sensitivity analysis that accounts for completion probability.",
        "Predeclare technical exclusions and response-time handling after the pilot. Log-transform heavily skewed response times instead of deleting valid slow responses without justification.",
    ]:
        bullet(doc, text)

    heading(doc, "6. Consent and privacy")
    paragraph(doc, "The deadline does not remove the need for school approval. For participants under 18, use the parent or guardian consent and student assent process required by the relevant school or ethics authority. The app's current checkbox saying that a guardian consented should be treated as an acknowledgement, not as proof of an approved consent process.")
    paragraph(doc, "Participation must not affect grades, attendance, teacher evaluation, or access to school activities. Keep the school-held invitation mapping separate from the research database. Define who can access exports, how long data will be retained, and how small groups will be protected in the final report.")

    heading(doc, "7. Minimum checklist before launch")
    for text in [
        "The student target population and allowed claims are written down.",
        "The school has approved recruitment, consent, session timing, and data fields.",
        "Random student invite lists and same-grade or same-section replacements exist.",
        "Teacher invitations are private and separate.",
        "Role, grade, revised age, invitation status, session conditions, app version, and dropout stage are recorded.",
        "The main app version and task order have been tested and frozen.",
        "The participant-level split and analysis code have been tested on synthetic data.",
        "The data cutoff, exclusion rules, target metrics, baselines, and claim limits are fixed before final analysis.",
    ]:
        bullet(doc, text)

    heading(doc, "Sources")
    sources = [
        ("ICMR: National Ethical Guidelines for Biomedical Research Involving Children", "https://www.icmr.gov.in/icmrobject/custom_data/pdf/resource-guidelines/National_Ethical_Guidelines_for_BioMedical_Research_Involving_Children_0.pdf"),
        ("STROBE Statement", "https://www.equator-network.org/reporting-guidelines/strobe/"),
        ("AAPOR Task Force on Non-Probability Sampling", "https://aapor.org/wp-content/uploads/2022/11/NPS_TF_Report_Final_7_revised_FNL_6_22_13.pdf"),
        ("Riley et al. (2020): Sample size for prediction model development", "https://www.bmj.com/content/368/bmj.m441"),
        ("Varoquaux (2018): Small samples and cross-validation error", "https://pubmed.ncbi.nlm.nih.gov/28655633/"),
        ("TRIPOD+AI reporting guidance", "https://pmc.ncbi.nlm.nih.gov/articles/PMC11019967/"),
        ("Bridges et al. (2020): Timing in online and laboratory experiments", "https://pubmed.ncbi.nlm.nih.gov/33005482/"),
        ("Sievertsen et al. (2016): Cognitive fatigue and student performance", "https://pmc.ncbi.nlm.nih.gov/articles/PMC4790980/"),
        ("Davidow et al. (2021): Probabilistic learning across ages 8-30", "https://pmc.ncbi.nlm.nih.gov/articles/PMC8279421/"),
        ("de Water et al. (2014): Age and adolescent delay discounting", "https://pmc.ncbi.nlm.nih.gov/articles/PMC3872775/"),
    ]
    for label, url in sources:
        p = doc.add_paragraph(style="List Bullet")
        add_hyperlink(p, label, url)

    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
        if p.style.name.startswith("Heading"):
            p.paragraph_format.keep_with_next = True

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
