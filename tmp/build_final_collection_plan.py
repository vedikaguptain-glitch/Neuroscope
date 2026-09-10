from datetime import date
from pathlib import Path

from docx import Document

from build_simple_protocol import (
    add_hyperlink,
    bullet,
    configure_styles,
    heading,
    numbered,
    paragraph,
    set_font,
)


ROOT = Path(r"C:\Users\USER\Desktop\code\Neuroscope")
OUT = ROOT / "output" / "documents" / "neuroscope-data-collection-plan-final.docx"


def add_title(doc):
    p = doc.add_paragraph()
    r = p.add_run("NEUROSCOPE Data Collection Plan")
    set_font(r, size=20, bold=True)

    p = doc.add_paragraph()
    r = p.add_run("Final plan for collecting student and teacher data before the science fair")
    set_font(r, italic=True)

    p = doc.add_paragraph()
    r = p.add_run("Prepared: ")
    set_font(r, bold=True)
    r = p.add_run(date.today().strftime("%d %B %Y"))
    set_font(r)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "NEUROSCOPE Data Collection Plan"
    doc.core_properties.author = "NEUROSCOPE research team"
    add_title(doc)

    heading(doc, "Plan in short")
    for text in [
        "The target is 100 completed student sessions. The minimum for the planned student analysis is 80 completed sessions.",
        "The 80 or 100 students must not include teachers. Teachers are a separate group.",
        "Randomly choose whole class sections from the participating grades. Invite every student in those sections who has the required permission.",
        "Use the normal app link during planned school sessions. Do not post it across the whole school.",
        "Invite every eligible teacher, but report teacher results separately.",
        "Use the simple model as the main model and the Transformer as a second, experimental comparison.",
        "Keep every participant completely inside one training or testing group.",
        "Limit the final claim to the selected classes in this school during the pre-exam period.",
    ]:
        bullet(doc, text)

    heading(doc, "1. Study goal")
    paragraph(doc, "The main study will examine students from selected classes in this school. The recommended prediction question is: Using a student's results from four activities, can a model predict whether that student chooses the gamble in at least half of the risk activity trials?")
    paragraph(doc, "A simple model should answer this question first. The Transformer should be compared with it using the same students and the same testing method. The Transformer should be presented as experimental unless it clearly and consistently performs better on students it did not train on.")
    paragraph(doc, "The 240 trials from one student provide detailed information about that student. They still count as one participant, not 240 participants.")

    heading(doc, "2. Sample size")
    paragraph(doc, "Aim for 100 completed student sessions. Do not stop at 100 invitations because some students may be absent, may not receive permission, may stop early, or may have technical problems. Plan to invite about 120 to 130 students so that roughly 100 can complete the study.")
    paragraph(doc, "The minimum is 80 completed student sessions. If fewer than 80 students complete the study by the deadline, report the app and model as a pilot and avoid strong performance claims. Teacher sessions cannot be added to the student count to reach 80.")
    bullet(doc, "Target: 100 completed students.")
    bullet(doc, "Minimum: 80 completed students.")
    bullet(doc, "Planned invitations: about 120 to 130 students.")
    bullet(doc, "Teachers: counted and analyzed separately.")

    heading(doc, "3. How to choose students")
    numbered(doc, "List the classes", "Write down every eligible section in each participating grade and the number of students in each section.")
    numbered(doc, "Choose sections before collection", "Randomly select enough sections to invite about 120 to 130 students. Spread the selected sections across the participating grades instead of choosing all students from one grade.")
    numbered(doc, "Choose backup sections", "Randomly choose one backup section for each grade before collection begins. Use a backup only if a selected section cannot take part or the total is still below 80 after the planned sessions.")
    numbered(doc, "Invite the full selected section", "Give the study information and permission form to every student in the selected section. Invite every student who returns the required permission and personally agrees to take part.")
    numbered(doc, "Use the normal link", "Share the ordinary app link in the selected class group shortly before the session or display it in the room. Do not share it in other class groups.")
    numbered(doc, "Record the numbers", "For each selected section, record how many students were enrolled, received the information, had permission, attended, started, completed, and stopped early. Keep this count sheet separate from anonymous app responses.")
    paragraph(doc, "Whole-class sampling is practical, but students from the same class may be similar. Include several sections across grades if possible. With only a few selected sections, describe the result as applying to those selected classes rather than the whole school.")

    heading(doc, "4. Teachers")
    paragraph(doc, "Invite all eligible teachers because there are fewer of them. Offer one or two staff-room sessions using the same app link. Keep teacher data out of student model training and out of the 80-student minimum.")
    paragraph(doc, "Report the number of teachers invited, started, and completed. Show simple teacher task summaries only when the group is large enough to protect individuals. A small teacher sample is an adult comparison, not proof that the model works for teachers generally.")

    heading(doc, "5. Session setup")
    for text in [
        "Run a pilot with 8 to 12 people before the main sessions. If the app or instructions change after the pilot, keep pilot data out of the main dataset.",
        "Reserve 40 to 45 minutes until the pilot gives a reliable duration. Make the duration shown in the app match the pilot result.",
        "Run student sessions in a quiet school room during approved zero periods, free periods, lunch, activity periods, or another low-disruption slot.",
        "Use the same type of device, browser, input method, instructions, task order, and app version as far as possible.",
        "Have one researcher read the same short introduction in every session. Teachers should not watch individual student answers.",
        "Offer one planned make-up session for students from selected classes who had permission but were absent.",
        "Treat reaction time as a secondary result if devices or browsers differ between sessions.",
    ]:
        bullet(doc, text)

    heading(doc, "6. Information the app must collect")
    for text in [
        "Role first: Student or Teacher.",
        "For students: grade and a school-approved class label.",
        "For students: completed age or narrower age choices. Wider age groups can be used in the report.",
        "For teachers: broad age groups and Prefer not to say. Do not ask for school grade.",
        "Collection setting: supervised school session, teacher session, or another approved setting.",
        "App version, session start time, session end time, and completion status saved to the database.",
        "A clear One submission per person message. Use the supervised attendance sheet to control repeats without connecting names to responses.",
        "No participant names in the research database.",
    ]:
        bullet(doc, text)

    heading(doc, "7. Age groups")
    paragraph(doc, "Ask whether the participant is a student or teacher before asking age questions. For students, collect completed age if the school approves it. In the report, combine ages into 13-14, 15-16, and 17-18 when each group is large enough to report safely.")
    paragraph(doc, "For teachers, broad groups such as 20-29, 30-39, 40-49, 50-59, and 60+ are enough. If the teacher group is very small, leave teacher age out of the report.")
    paragraph(doc, "Do not claim that age caused a difference between students and teachers. In one school, age and role mostly overlap, so their effects cannot be separated properly.")

    heading(doc, "8. Permission and privacy")
    paragraph(doc, "Get written school approval before collecting the main data. For students under 18, follow the parent or guardian permission process required by the school or ethics authority. The app should ask the student for their own agreement to participate; it should not ask the student to confirm that a guardian has already consented.")
    paragraph(doc, "Participation must be voluntary. Refusing or stopping must not affect grades, attendance, or treatment by teachers. Explain the purpose, expected length, recorded information, possible tiredness, right to stop, data access, storage period, and contact person in simple language.")
    paragraph(doc, "The school may keep a separate permission and attendance list. Do not join that list to anonymous participant responses. Decide before collection who can open the research data and when it will be deleted.")

    heading(doc, "9. Training and testing with 80 to 100 students")
    paragraph(doc, "A single small testing group can give a lucky or unlucky score. With 80 to 100 students, use repeated five-part testing so that every student is tested while still keeping training and testing separate.")
    bullet(doc, "Create five groups. Keep grades and the two target answers as balanced as possible.")
    bullet(doc, "Train on four groups. Choose model settings using only those training students.")
    bullet(doc, "Test on the remaining group. That group must remain unseen during the round.")
    bullet(doc, "Rotate the test group until each group has served as the unseen group.")
    bullet(doc, "Repeat the full five-part process with several saved random splits. Report the average result and how much it changes.")
    paragraph(doc, "All trials from one participant must stay together. Response-time averages, scaling, vocabularies, missing-data rules, feature selection, encoder training, early stopping, and model settings must be fitted using training participants only.")
    paragraph(doc, "The current model evaluation code must train the encoder separately inside each outer training split. A score from an encoder trained on the full dataset must not be used as the final science-fair result.")

    heading(doc, "10. What to calculate")
    for text in [
        "Recruitment: enrolled, permitted, started, completed, and stopped counts for every selected section and for teachers.",
        "Data quality: missing trials, early stops, session duration, device or browser group, and repeated or suspicious sessions found during supervised review.",
        "Behavior: the main choice rate for each task and the middle response time, which is less affected by a few extremely slow responses.",
        "Prediction: results from the simple model, the Transformer, and the rule that always predicts the more common answer.",
        "Scores: confusion matrix, balanced accuracy, and a 95% range made by resampling whole participants.",
        "Stability: the average result and variation across the repeated participant-level splits.",
        "Class differences: results by selected section or grade. Do not hide large differences between classes inside one overall score.",
        "Partial sessions: the number and point where people stopped. Do not silently delete every session with fewer than 240 trials.",
    ]:
        bullet(doc, text)

    heading(doc, "11. Timeline")
    bullet(doc, "Day 1: request school approval, list eligible sections, and randomly choose the main and backup sections.")
    bullet(doc, "Days 1 to 3: make the minimum app changes, test the database export, and prepare the parent information, student agreement, and standard session script.")
    bullet(doc, "Days 3 and 4: run the 8 to 12 person pilot. Fix serious problems and then freeze the app and analysis plan.")
    bullet(doc, "Days 5 to 12: run selected class sessions, the make-up session, and separate teacher sessions.")
    bullet(doc, "If fewer than 80 students have completed, use the preselected backup sections while time and approval allow. Continue toward the target of 100.")
    bullet(doc, "Final days: close the dataset, save an untouched copy, check completion and repeats, run the fixed analysis, and prepare the fair report.")
    paragraph(doc, "If the school does not approve student collection in time, do not circulate the link informally to minors. Submit the working app, any properly approved pilot or teacher results, the tested analysis process using demonstration data, and the future collection plan.")

    heading(doc, "12. Claims for the science fair")
    bullet(doc, "Allowed: We collected data from at least 80 students in randomly selected classes and tested whether behavior in four activities predicted risk choices.")
    bullet(doc, "Allowed: We tested unseen students through repeated participant-level splits and reported how much the result changed.")
    bullet(doc, "Allowed: Teacher results are a small, separate description.")
    bullet(doc, "Not allowed: The model works for every student, teacher, age group, school, or person.")
    bullet(doc, "Not allowed: The number of trials replaces the need for enough participants.")
    bullet(doc, "Not allowed: A strong score from one split proves that the model is reliable.")

    heading(doc, "13. Final checklist")
    for text in [
        "The target is 100 completed students and the minimum is 80.",
        "About 120 to 130 students will be invited across randomly selected sections.",
        "Backup sections were chosen before collection.",
        "Teachers are separate from the student target and analysis.",
        "The school approved the study, schedule, permission process, and data fields.",
        "The app records role, student grade or teacher status, revised age, setting, app version, and server-side completion.",
        "The normal link will be used only in selected class and teacher sessions.",
        "The app duration matches the pilot result.",
        "Every participant stays inside one training or testing group.",
        "All data preparation and model training happen inside the training part of each split.",
        "The report includes recruitment counts, partial sessions, uncertainty, and clear limits.",
    ]:
        bullet(doc, text)

    heading(doc, "Sources")
    sources = [
        ("ICMR: ethical guidance for research involving children", "https://www.icmr.gov.in/icmrobject/custom_data/pdf/resource-guidelines/National_Ethical_Guidelines_for_BioMedical_Research_Involving_Children_0.pdf"),
        ("CDC: school surveys that select whole classes", "https://www.cdc.gov/mmwr/preview/mmwrhtml/rr5312a1.htm"),
        ("AAPOR: selection and inference with non-probability samples", "https://aapor.org/wp-content/uploads/2023/02/Task-Force-Report-FINAL.pdf"),
        ("TRIPOD+AI: reporting model development and evaluation", "https://www.bmj.com/content/385/bmj-2023-078378"),
        ("Scikit-learn: avoiding leakage and using nested testing", "https://scikit-learn.org/stable/auto_examples/model_selection/plot_nested_cross_validation_iris.html"),
        ("Bridges and colleagues: browser and device timing differences", "https://pmc.ncbi.nlm.nih.gov/articles/PMC7512138/"),
    ]
    for label, url in sources:
        p = doc.add_paragraph()
        add_hyperlink(p, label, url)

    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
        if p.style.name.startswith("Heading"):
            p.paragraph_format.keep_with_next = True

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
